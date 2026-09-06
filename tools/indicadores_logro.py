"""Lectura y escritura de `TablaIndicadoresLogro` (hoja IndicadoresLogro).

Misma filosofía que `situaciones_aprendizaje`: leer con openpyxl (inofensivo),
escribir de forma quirúrgica sobre el ZIP para no cargarse el modelo de Power
Pivot ni las tablas dinámicas.

Estructura de `TablaIndicadoresLogro` (rango A4:M{n+4}), columnas:

    CE   criterio de evaluación (p.ej. "4.2")   -> se elige de TablaCriteriosEvaluacion
    CED  descripción del criterio               -> automático (lookup por CE)
    IL   indicador de logro ("4.2.1", "4.2.2")  -> automático, correlativo dentro del CE
    PIL  peso del indicador                     -> a mano (número)
    PIL% PIL / Σ PIL del mismo CE               -> automático (fórmula en el Excel)
    DIL  descripción del indicador              -> a mano
    DO   descriptores operativos                -> a mano
    CON  contenidos                             -> a mano
    CT   (carga temporal / referencias)         -> a mano
    IE   instrumento de evaluación              -> lista: TablasAuxiliares!B2:B11
    CC   criterio de calificación              -> lista: TablasAuxiliares!F2:F5
    AE   agente evaluador                       -> lista: TablasAuxiliares!D2:D4
    SA   situación de aprendizaje               -> lista: SA existentes en tablaSAprendizaje

La fila de cabecera está en la fila 4 (filas 1-3 llevan un rótulo suelto).
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import load_workbook

SHEET_NAME = "IndicadoresLogro"
TABLE_NAME = "TablaIndicadoresLogro"
CRITERIOS_SHEET_PREFIX = "CriteriosEvaluaci"
AUX_SHEET = "TablasAuxiliares"

COLUMNS = ["CE", "CED", "IL", "PIL", "PIL%", "DIL", "DO", "CON", "CT", "IE", "CC", "AE", "SA"]
COL_LETTER = {name: chr(ord("A") + i) for i, name in enumerate(COLUMNS)}
HEADER_ROW = 4
FIRST_DATA_ROW = 5

TEXT_COLS = {"CE", "CED", "IL", "DIL", "DO", "CON", "CT", "IE", "CC", "AE"}
NUM_COLS = {"PIL", "SA"}
AUTO_COLS = {"CED", "IL", "PIL%"}  # el usuario no las teclea

# Estilos de celda por columna observados en el Excel original.
CELL_STYLE = {"CED": "1", "PIL%": "5", "DIL": "1", "DO": "1", "CON": "2", "CT": "2"}

PIL_PCT_FORMULA = (
    "{t}[[#This Row],[PIL]]/SUMIF({t}[CE], {t}[[#This Row],[CE]], {t}[PIL])"
).format(t=TABLE_NAME)

AUX_RANGES = {  # columna del indicador -> (columna en TablasAuxiliares, filas)
    "IE": ("B", range(2, 12)),
    "CC": ("F", range(2, 6)),
    "AE": ("D", range(2, 5)),
}


def _read_source_bytes(source: Any) -> bytes:
    if isinstance(source, (str, Path)):
        return Path(source).expanduser().resolve().read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    if hasattr(source, "read"):
        return source.read()
    raise TypeError("El origen debe ser una ruta o un archivo subido")


def _find_sheet(wb, name_or_prefix: str, *, prefix: bool = False):
    for name in wb.sheetnames:
        if (name.startswith(name_or_prefix) if prefix else name == name_or_prefix):
            return wb[name]
    return None


def _s(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _num(value: Any):
    if value in (None, ""):
        return None
    try:
        f = float(value)
    except (TypeError, ValueError):
        return None
    return int(f) if f.is_integer() else f


def read_indicadores(source: Any) -> dict:
    payload = _read_source_bytes(source)
    wb = load_workbook(BytesIO(payload), data_only=True)

    ws = _find_sheet(wb, SHEET_NAME)
    if ws is None:
        raise ValueError(f"El Excel no tiene una hoja '{SHEET_NAME}'.")

    rows: list[dict] = []
    for r in ws.iter_rows(min_row=FIRST_DATA_ROW, max_col=len(COLUMNS), values_only=True):
        if r is None or all(c in (None, "") for c in r):
            continue
        row = {COLUMNS[i]: (r[i] if i < len(r) else None) for i in range(len(COLUMNS))}
        rows.append(
            {
                "CE": _s(row["CE"]),
                "CED": _s(row["CED"]),
                "IL": _s(row["IL"]),
                "PIL": _num(row["PIL"]),
                "DIL": _s(row["DIL"]),
                "DO": _s(row["DO"]),
                "CON": _s(row["CON"]),
                "CT": _s(row["CT"]),
                "IE": _s(row["IE"]),
                "CC": _s(row["CC"]),
                "AE": _s(row["AE"]),
                "SA": _num(row["SA"]),
            }
        )

    # CE -> CED y peso P desde TablaCriteriosEvaluacion
    ce_ced: dict[str, str] = {}
    ce_p: dict[str, float] = {}
    ce_list: list[str] = []
    ws_ce = _find_sheet(wb, CRITERIOS_SHEET_PREFIX, prefix=True)
    if ws_ce is not None:
        seen = False
        for r in ws_ce.iter_rows(values_only=True):
            if not r:
                continue
            if not seen:
                if r[0] is not None and str(r[0]).strip() == "CE":
                    seen = True
                continue
            ce = _s(r[0])
            if not ce:
                continue
            ce_ced[ce] = _s(r[1]) if len(r) > 1 else ""
            ce_p[ce] = _num(r[2]) or 0.0 if len(r) > 2 else 0.0
            ce_list.append(ce)

    # Listas auxiliares
    aux: dict[str, list[str]] = {}
    ws_aux = _find_sheet(wb, AUX_SHEET)
    if ws_aux is not None:
        for key, (col, filas) in AUX_RANGES.items():
            vals = []
            for rr in filas:
                v = ws_aux[f"{col}{rr}"].value
                if v not in (None, ""):
                    vals.append(str(v).strip())
            aux[key] = vals

    return {
        "rows": rows,
        "ce_list": ce_list,
        "ce_ced": ce_ced,
        "ce_p": ce_p,
        "aux": aux,
    }


def _split_tokens(value: Any) -> list[str]:
    """Divide "A.1.1, A 1 2 ,B.3.1" en ['A.1.1', 'A.1.2', 'B.3.1'] (espacios
    internos -> puntos, como hace la hoja RESUMEN con SUBSTITUTE)."""
    out = []
    for tok in str(value or "").split(","):
        tok = re.sub(r"\s+", ".", tok.strip()).strip(".")
        if tok:
            out.append(tok)
    return out


def matriz_sa_ce(
    rows: list[dict],
    ce_list: list[str],
    ce_p: dict[str, float] | None = None,
) -> dict:
    """Réplica de la tabla dinámica "DISTRIBUCIÓN DE PORCENTAJES POR CRITERIOS DE
    EVALUACIÓN Y SITUACIONES DE APRENDIZAJE" de la hoja INFORMES.

    Para cada indicador de logro, su peso sobre TODA la programación es
    ``(PIL / ΣPIL del CE) × (P del CE / ΣP)``. La matriz agrupa esos pesos por
    (SA, CE). Los totales por fila dan el peso de cada SA según los IL asignados.
    """
    ce_p = ce_p or {}
    total_p = sum(ce_p.values())

    sum_pil: dict[str, float] = {}
    for r in rows:
        ce = _s(r.get("CE"))
        pil = r.get("PIL")
        pil = 0.0 if pil in (None, "") else float(pil)
        sum_pil[ce] = sum_pil.get(ce, 0.0) + pil

    matrix: dict[tuple[int, str], float] = {}
    sa_set: set[int] = set()
    for r in rows:
        ce = _s(r.get("CE"))
        sa = r.get("SA")
        if sa in (None, "") or not ce:
            continue
        sa = int(sa)
        pil = r.get("PIL")
        pil = 0.0 if pil in (None, "") else float(pil)
        spil = sum_pil.get(ce, 0.0)
        pce = ce_p.get(ce, 0.0)
        w = (pil / spil if spil else 0.0) * (pce / total_p if total_p else 0.0)
        matrix[(sa, ce)] = matrix.get((sa, ce), 0.0) + w
        sa_set.add(sa)

    sa_rows = sorted(sa_set)
    ce_cols = [c for c in ce_list if c in sum_pil] or list(sum_pil.keys())
    sa_total = {sa: sum(matrix.get((sa, c), 0.0) for c in ce_cols) for sa in sa_rows}
    ce_total = {c: sum(matrix.get((sa, c), 0.0) for sa in sa_rows) for c in ce_cols}
    grand = sum(sa_total.values())
    return {
        "matrix": matrix,
        "sa_rows": sa_rows,
        "ce_cols": ce_cols,
        "sa_total": sa_total,
        "ce_total": ce_total,
        "grand_total": grand,
    }


def resumen_por_ce(
    rows: list[dict],
    ce_list: list[str],
    ce_ced: dict[str, str],
    ce_p: dict[str, float] | None = None,
) -> list[dict]:
    """Resumen por criterio de evaluación, equivalente a la hoja RESUMEN:
    por cada CE, su %CE (P/ΣP), y la lista única de contenidos (CON), CT y SA
    de todos sus indicadores de logro."""
    ce_p = ce_p or {}
    total_p = sum(ce_p.values())

    por_ce: dict[str, dict] = {}
    orden: list[str] = []
    for r in rows:
        ce = _s(r.get("CE"))
        if not ce:
            continue
        if ce not in por_ce:
            por_ce[ce] = {"CON": [], "CT": [], "SA": [], "IL": []}
            orden.append(ce)
        d = por_ce[ce]
        d["CON"] += _split_tokens(r.get("CON"))
        d["CT"] += _split_tokens(r.get("CT"))
        if r.get("SA") not in (None, ""):
            d["SA"].append(str(int(r["SA"])) if isinstance(r["SA"], (int, float)) else str(r["SA"]))
        if _s(r.get("IL")):
            d["IL"].append(_s(r["IL"]))

    secuencia = [c for c in ce_list if c in por_ce] + [c for c in orden if c not in ce_list]

    def _uniq(seq):
        return list(dict.fromkeys(seq))

    out = []
    for ce in secuencia:
        d = por_ce[ce]
        p = ce_p.get(ce, 0.0)
        out.append(
            {
                "CE": ce,
                "CED": ce_ced.get(ce, ""),
                "pct": (p / total_p) if total_p else 0.0,
                "IL": ", ".join(_uniq(d["IL"])),
                "CONTENIDOS": ", ".join(_uniq(d["CON"])),
                "CT": ", ".join(_uniq(d["CT"])),
                "SA": ", ".join(_uniq(d["SA"])),
            }
        )
    return out


def recompute(
    rows: list[dict],
    ce_ced: dict[str, str],
    ce_order: list[str] | None = None,
) -> list[dict]:
    """Agrupa las filas por CE (en el orden de `ce_order`, o de aparición si no se
    pasa), rellena CED (lookup por CE), numera IL correlativo dentro de cada CE
    (4.2.1, 4.2.2…) y calcula PIL% = PIL / ΣPIL del CE.

    Al agrupar por CE, una fila nueva de "3.2" se coloca junto a las demás de
    "3.2" y recibe el siguiente número (3.2.4, 3.2.5…); dentro de un mismo CE se
    respeta el orden en que estaban las filas."""
    limpio = [r for r in rows if _row_has_content(r)]

    orden = list(ce_order or [])
    rank_cache: dict[str, int] = {}

    def _rank(ce: str) -> int:
        if ce in rank_cache:
            return rank_cache[ce]
        if ce in orden:
            rank_cache[ce] = orden.index(ce)
        elif ce:
            rank_cache[ce] = len(orden) + len(rank_cache)
        else:
            rank_cache[ce] = 10**9  # filas sin CE, al final
        return rank_cache[ce]

    # sorted() es estable: dentro de un mismo CE se mantiene el orden previo.
    limpio = sorted(limpio, key=lambda r: _rank(_s(r.get("CE"))))

    contador: dict[str, int] = {}
    suma_pil: dict[str, float] = {}
    for r in limpio:
        ce = _s(r.get("CE"))
        pil = r.get("PIL")
        pil = 1 if pil in (None, "") else pil
        try:
            suma_pil[ce] = suma_pil.get(ce, 0.0) + float(pil)
        except (TypeError, ValueError):
            pass

    out: list[dict] = []
    for r in limpio:
        ce = _s(r.get("CE"))
        pil = r.get("PIL")
        pil = 1 if pil in (None, "") else pil
        contador[ce] = contador.get(ce, 0) + 1
        try:
            pil_f = float(pil)
        except (TypeError, ValueError):
            pil_f = 0.0
        total = suma_pil.get(ce, 0.0)
        out.append(
            {
                **r,
                "CE": ce,
                "CED": ce_ced.get(ce, _s(r.get("CED"))),
                "IL": f"{ce}.{contador[ce]}" if ce else "",
                "PIL": pil,
                "PIL%": (pil_f / total) if total else 0.0,
            }
        )
    return out


def _row_has_content(r: dict) -> bool:
    return any(
        _s(r.get(c)) for c in ("CE", "DIL", "DO", "CON", "CT", "IE", "CC", "AE")
    ) or r.get("PIL") not in (None, "") or r.get("SA") not in (None, "")


# ---------------------------------------------------------------------------
# ESCRITURA QUIRÚRGICA
# ---------------------------------------------------------------------------
def _find_sheet_path(zf: zipfile.ZipFile, sheet_name: str) -> str:
    wb = zf.read("xl/workbook.xml").decode("utf-8")
    m = re.search(rf'<sheet[^>]*name="{re.escape(sheet_name)}"[^>]*r:id="([^"]+)"', wb)
    rid = m.group(1)
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    target = re.search(
        rf'<Relationship[^>]*Id="{re.escape(rid)}"[^>]*Target="([^"]+)"', rels
    ).group(1).lstrip("/")
    return target if target.startswith("xl/") else "xl/" + target.replace("../", "")


def _find_table_path(zf: zipfile.ZipFile, table_name: str) -> str | None:
    for name in zf.namelist():
        if name.startswith("xl/tables/") and name.endswith(".xml"):
            if f'name="{table_name}"' in zf.read(name).decode("utf-8"):
                return name
    return None


def _num_xml(value) -> str:
    f = float(value)
    return str(int(f)) if f.is_integer() else repr(f)


def _cell(ref: str, col: str, value) -> str:
    style = f' s="{CELL_STYLE[col]}"' if col in CELL_STYLE else ""
    if col == "PIL%":
        return f'<c r="{ref}"{style}><f>{escape(PIL_PCT_FORMULA)}</f></c>'
    if value in (None, ""):
        return f'<c r="{ref}"{style}/>' if style else ""
    if col in NUM_COLS:
        n = _num(value)
        return f'<c r="{ref}"{style}><v>{_num_xml(n)}</v></c>' if n is not None else ""
    return (
        f'<c r="{ref}"{style} t="inlineStr"><is>'
        f'<t xml:space="preserve">{escape(str(value))}</t></is></c>'
    )


def _rewrite_sheet(sheet_xml: str, rows: list[dict]) -> str:
    sd = re.search(r"<sheetData>(.*)</sheetData>", sheet_xml, re.S)
    rows_xml = re.findall(r"<row [^>]*>.*?</row>|<row [^>]*/>", sd.group(1), re.S)
    by_num = {int(re.search(r'r="(\d+)"', rx).group(1)): rx for rx in rows_xml}

    keep = [by_num[n] for n in sorted(by_num) if n <= HEADER_ROW]  # filas 1..4 intactas

    n = len(rows)
    for i, r in enumerate(rows):
        rn = FIRST_DATA_ROW + i
        cells = "".join(
            _cell(f"{COL_LETTER[c]}{rn}", c, r.get(c)) for c in COLUMNS
        )
        keep.append(f'<row r="{rn}" spans="1:13" x14ac:dyDescent="0.25">{cells}</row>')

    last_row = HEADER_ROW + max(n, 1)
    sheet_xml = (
        sheet_xml[: sd.start()] + "<sheetData>" + "".join(keep) + "</sheetData>" + sheet_xml[sd.end():]
    )
    sheet_xml = re.sub(
        r'<dimension ref="[^"]*"/>', f'<dimension ref="A1:M{last_row}"/>', sheet_xml
    )
    # Ajusta los rangos de las listas desplegables (IE/CC/AE) a las filas de datos.
    for col in ("J", "K", "L"):
        sheet_xml = re.sub(
            rf"<xm:sqref>{col}\d+:{col}\d+</xm:sqref>",
            f"<xm:sqref>{col}{FIRST_DATA_ROW}:{col}{last_row}</xm:sqref>",
            sheet_xml,
        )
    return sheet_xml


def _rewrite_table(table_xml: str, n: int) -> str:
    ref = f"A{HEADER_ROW}:M{HEADER_ROW + max(n, 1)}"
    table_xml = re.sub(r'(<table[^>]*\sref=")[^"]*(")', rf"\g<1>{ref}\g<2>", table_xml)
    table_xml = re.sub(r'(<autoFilter\s+ref=")[^"]*(")', rf"\g<1>{ref}\g<2>", table_xml)
    return table_xml


def _force_full_recalc(workbook_xml: str) -> str:
    if "fullCalcOnLoad" in workbook_xml:
        return workbook_xml
    if "<calcPr" in workbook_xml:
        return re.sub(r"<calcPr\s+", '<calcPr fullCalcOnLoad="1" ', workbook_xml, count=1)
    return workbook_xml.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')


def save_indicadores(
    source: Any,
    rows: list[dict],
    ce_ced: dict[str, str],
    ce_order: list[str] | None = None,
) -> bytes:
    payload = _read_source_bytes(source)
    rows = recompute(rows, ce_ced, ce_order)

    src = zipfile.ZipFile(BytesIO(payload))
    sheet_path = _find_sheet_path(src, SHEET_NAME)
    table_path = _find_table_path(src, TABLE_NAME)

    replacements = {
        sheet_path: _rewrite_sheet(src.read(sheet_path).decode("utf-8"), rows).encode("utf-8"),
        "xl/workbook.xml": _force_full_recalc(
            src.read("xl/workbook.xml").decode("utf-8")
        ).encode("utf-8"),
    }
    if table_path:
        replacements[table_path] = _rewrite_table(
            src.read(table_path).decode("utf-8"), len(rows)
        ).encode("utf-8")

    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, replacements.get(item.filename, src.read(item.filename)))
    return out.getvalue()


CRITERIOS_TABLE = "TablaCriteriosEvaluacion"


def _find_sheet_path_prefix(zf: zipfile.ZipFile, prefix: str) -> str:
    wb = zf.read("xl/workbook.xml").decode("utf-8")
    m = re.search(rf'<sheet[^>]*name="({re.escape(prefix)}[^"]*)"[^>]*r:id="([^"]+)"', wb)
    rid = m.group(2)
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    target = re.search(
        rf'<Relationship[^>]*Id="{re.escape(rid)}"[^>]*Target="([^"]+)"', rels
    ).group(1).lstrip("/")
    return target if target.startswith("xl/") else "xl/" + target.replace("../", "")


def save_criterios(source: Any, ce_p: dict[str, float], ce_order: list[str]) -> bytes:
    """Reescribe solo la columna de pesos P de `TablaCriteriosEvaluacion`
    (mismo nº de filas: aquí no se añaden ni quitan criterios)."""
    payload = _read_source_bytes(source)
    src = zipfile.ZipFile(BytesIO(payload))

    sheet_path = _find_sheet_path_prefix(src, CRITERIOS_SHEET_PREFIX)
    table_path = _find_table_path(src, CRITERIOS_TABLE)
    sheet_xml = src.read(sheet_path).decode("utf-8")

    first_row = 5
    if table_path:
        ref = re.search(r'<table[^>]*\sref="[A-Z]+(\d+):', src.read(table_path).decode("utf-8"))
        if ref:
            first_row = int(ref.group(1)) + 1

    for i, ce in enumerate(ce_order):
        if ce not in ce_p:
            continue
        row = first_row + i
        val = _num_xml(ce_p[ce])
        sheet_xml = re.sub(
            rf'<c r="C{row}"[^>]*>(?:<v>[^<]*</v>)?</c>',
            f'<c r="C{row}"><v>{val}</v></c>',
            sheet_xml,
            count=1,
        )
        # quita el valor en caché del P% de esa fila
        sheet_xml = re.sub(
            rf'(<c r="D{row}"[^>]*>(?:<f[^>]*>.*?</f>|<f[^>]*/>))<v>[^<]*</v>',
            r"\g<1>",
            sheet_xml,
        )

    replacements = {
        sheet_path: sheet_xml.encode("utf-8"),
        "xl/workbook.xml": _force_full_recalc(
            src.read("xl/workbook.xml").decode("utf-8")
        ).encode("utf-8"),
    }
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, replacements.get(item.filename, src.read(item.filename)))
    return out.getvalue()
