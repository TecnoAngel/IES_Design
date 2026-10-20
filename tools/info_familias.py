"""Documento de información a las familias sobre la evaluación (ESO / Bachillerato).

Responde a la obligación de informar al alumnado y a las familias (ORDEN
EDU/424/2024, art. 12 y ORDEN EDU/425/2024, art. 10): competencias específicas,
criterios de evaluación, contenidos y su temporalización, procedimientos e
instrumentos de evaluación, criterios de calificación (peso de los criterios y
redondeo), criterios de promoción/titulación y medidas de atención
individualizada. Además, una página por situación de aprendizaje con los
criterios e indicadores de logro que se trabajan (sin el texto del indicador),
sus pesos y la tarta de instrumentos de esa SA.

Lo que no está en el resto del Excel (texto de las competencias específicas,
criterios de redondeo, promoción, medidas de atención…) se guarda en una hoja
propia ``Info_Familias``:

* ``A:B`` pares campo / valor (ver ``CAMPOS``);
* ``D:E`` nº de competencia específica / descripción.

La hoja no es una tabla de Excel (ListObject) y es enteramente de la app: al
guardar se reescribe entera. Si el Excel no la tiene, se **añade** de forma
quirúrgica sobre el ZIP (workbook.xml, sus rels, [Content_Types].xml y
docProps/app.xml), sin tocar el resto del libro.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import jinja2
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docxtpl import DocxTemplate, InlineImage, RichText
from openpyxl import load_workbook

from tools.programacion_aula import _set_cell_background, _set_table_borders

SHEET_NAME = "Info_Familias"

# (clave, etiqueta para la app, ayuda). El orden es el de la hoja.
CAMPOS: list[tuple[str, str, str]] = [
    ("redondeo_criterios", "Criterio de redondeo en la calificación de cada criterio de evaluación",
     "Según la Propuesta Curricular y la programación didáctica."),
    ("redondeo_materia", "Criterio de redondeo en la calificación de la materia o ámbito",
     "Según la Propuesta Curricular y la programación didáctica."),
    ("competencias_clave", "Calificación de las competencias clave",
     "Procedimiento y criterios de calificación de las competencias clave a partir de los "
     "criterios de evaluación, peso de los criterios y redondeo (Propuesta Curricular). "
     "En ESO, procedimiento para valorar la promoción excepcional."),
    ("promocion_titulacion", "Criterios de promoción y titulación",
     "Los establecidos en la normativa; los comunica la tutoría."),
    ("plan_recuperacion", "Plan de recuperación",
     "Para el alumnado con la materia pendiente de cursos anteriores."),
    ("plan_refuerzo", "Plan de refuerzo", "Para el alumnado que no promociona o con dificultades."),
    ("apoyos", "Apoyos individualizados o en grupo específico",
     "Apoyos fuera del grupo ordinario."),
    ("adaptaciones", "Adaptaciones curriculares", "Significativas o no significativas."),
]
CAMPOS_KEYS = [k for k, _l, _h in CAMPOS]

AGENTES = {"HE": "heteroevaluación (profesorado)", "CO": "coevaluación", "AEV": "autoevaluación"}

SIN_COMPLETAR = "(pendiente de completar)"
HEADER_FILL = "A6182E"
_JINJA_ENV = jinja2.Environment(undefined=jinja2.ChainableUndefined)

_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_REL_WORKSHEET = _NS_REL + "/worksheet"
_CT_WORKSHEET = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"


def _read_bytes(source: Any) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if isinstance(source, (str, Path)):
        return Path(source).expanduser().resolve().read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    if hasattr(source, "read"):
        return source.read()
    raise TypeError("El origen debe ser una ruta o un archivo subido")


def _s(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _fmt_pct(value: float) -> str:
    return f"{value * 100:.2f}".replace(".", ",") + " %"


def numero_competencia(ce: str) -> str:
    """'2.4' -> '2' (la competencia específica es el primer número del CE)."""
    return _s(ce).split(".", 1)[0]


# ---------------------------------------------------------------------------
# LECTURA
# ---------------------------------------------------------------------------
def read_info_familias(source: Any) -> dict:
    """{"campos": {clave: texto}, "competencias": {nº: descripción}, "existe":
    bool}. Vacío si el Excel aún no tiene la hoja."""
    out = {"campos": {k: "" for k in CAMPOS_KEYS}, "competencias": {}, "existe": False}
    wb = load_workbook(BytesIO(_read_bytes(source)), data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        return out
    out["existe"] = True
    ws = wb[SHEET_NAME]
    for row in ws.iter_rows(min_row=2, max_col=5, values_only=True):
        a, b = (row[0] if len(row) > 0 else None), (row[1] if len(row) > 1 else None)
        d, e = (row[3] if len(row) > 3 else None), (row[4] if len(row) > 4 else None)
        if _s(a) in out["campos"]:
            out["campos"][_s(a)] = _s(b)
        if _s(d):
            num = str(int(d)) if isinstance(d, float) and d.is_integer() else _s(d)
            out["competencias"][num] = _s(e)
    return out


# ---------------------------------------------------------------------------
# ESCRITURA QUIRÚRGICA (hoja propia, se reescribe entera; se añade si falta)
# ---------------------------------------------------------------------------
def _txt(ref: str, value: Any) -> str:
    if value in (None, ""):
        return ""
    return (
        f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">'
        f"{escape(str(value))}</t></is></c>"
    )


def _comp_sort_key(num: str):
    return (0, int(num)) if num.isdigit() else (1, num)


def _sheet_xml(campos: dict[str, str], competencias: dict[str, str]) -> str:
    comps = sorted(competencias.items(), key=lambda kv: _comp_sort_key(kv[0]))
    n = max(len(CAMPOS_KEYS), len(comps))
    rows = [
        '<row r="1">' + _txt("A1", "Campo") + _txt("B1", "Valor")
        + _txt("D1", "CEsp") + _txt("E1", "Competencia específica") + "</row>"
    ]
    for i in range(n):
        r = 2 + i
        cells = ""
        if i < len(CAMPOS_KEYS):
            k = CAMPOS_KEYS[i]
            cells += _txt(f"A{r}", k) + _txt(f"B{r}", campos.get(k, ""))
        if i < len(comps):
            cells += _txt(f"D{r}", comps[i][0]) + _txt(f"E{r}", comps[i][1])
        rows.append(f'<row r="{r}">{cells}</row>')
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<worksheet xmlns="{_NS_MAIN}" xmlns:r="{_NS_REL}">'
        f'<dimension ref="A1:E{n + 1}"/>'
        '<sheetViews><sheetView workbookViewId="0"/></sheetViews>'
        '<sheetFormatPr defaultRowHeight="15"/>'
        '<cols><col min="1" max="1" width="26" customWidth="1"/>'
        '<col min="2" max="2" width="90" customWidth="1"/>'
        '<col min="4" max="4" width="7" customWidth="1"/>'
        '<col min="5" max="5" width="90" customWidth="1"/></cols>'
        f'<sheetData>{"".join(rows)}</sheetData>'
        '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
        "</worksheet>"
    )


def _sheet_path(zf: zipfile.ZipFile) -> str | None:
    wb = zf.read("xl/workbook.xml").decode("utf-8")
    m = re.search(rf'<sheet[^>]*name="{re.escape(SHEET_NAME)}"[^>]*r:id="([^"]+)"', wb)
    if not m:
        return None
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    t = re.search(rf'<Relationship[^>]*Id="{re.escape(m.group(1))}"[^>]*Target="([^"]+)"', rels)
    tgt = t.group(1).lstrip("/")
    return tgt if tgt.startswith("xl/") else "xl/" + tgt.replace("../", "")


def _registrar_hoja(zf: zipfile.ZipFile) -> tuple[str, dict[str, bytes]]:
    """Da de alta una hoja nueva al final del libro. Devuelve (ruta de la hoja,
    partes del paquete modificadas)."""
    names = set(zf.namelist())
    n = 1
    while f"xl/worksheets/sheet{n}.xml" in names:
        n += 1
    path = f"xl/worksheets/sheet{n}.xml"

    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid_n = max([int(x) for x in re.findall(r'Id="rId(\d+)"', rels)] or [0]) + 1
    rid = f"rId{rid_n}"
    rels = rels.replace(
        "</Relationships>",
        f'<Relationship Id="{rid}" Type="{_REL_WORKSHEET}" Target="worksheets/sheet{n}.xml"/>'
        "</Relationships>",
    )

    wb = zf.read("xl/workbook.xml").decode("utf-8")
    sheet_id = max([int(x) for x in re.findall(r'<sheet [^>]*sheetId="(\d+)"', wb)] or [0]) + 1
    wb = wb.replace(
        "</sheets>", f'<sheet name="{SHEET_NAME}" sheetId="{sheet_id}" r:id="{rid}"/></sheets>', 1
    )

    ct = zf.read("[Content_Types].xml").decode("utf-8")
    ct = ct.replace(
        "</Types>", f'<Override PartName="/{path}" ContentType="{_CT_WORKSHEET}"/></Types>'
    )

    parts = {
        "xl/_rels/workbook.xml.rels": rels.encode("utf-8"),
        "xl/workbook.xml": wb.encode("utf-8"),
        "[Content_Types].xml": ct.encode("utf-8"),
    }

    # docProps/app.xml: nº de hojas y su nombre en TitlesOfParts (lo usa el
    # Explorador de Windows; Excel no lo exige, así que si no cuadra se deja).
    if "docProps/app.xml" in names:
        try:
            app = zf.read("docProps/app.xml").decode("utf-8")
            n_hojas = len(re.findall(r"<sheet ", wb)) - 1  # sin contar la nueva
            m = re.search(r"(<HeadingPairs>.*?<vt:i4>)(\d+)(</vt:i4>)", app, re.S)
            tp = re.search(r'(<TitlesOfParts><vt:vector size=")(\d+)(" baseType="lpstr">)(.*?)(</vt:vector>)', app, re.S)
            if m and tp and int(m.group(2)) == n_hojas:
                app = app[: m.start(2)] + str(n_hojas + 1) + app[m.end(2):]
                tp = re.search(r'(<TitlesOfParts><vt:vector size=")(\d+)(" baseType="lpstr">)(.*?)(</vt:vector>)', app, re.S)
                items = re.findall(r"<vt:lpstr>.*?</vt:lpstr>", tp.group(4), re.S)
                items.insert(n_hojas, f"<vt:lpstr>{SHEET_NAME}</vt:lpstr>")
                app = (
                    app[: tp.start()] + tp.group(1) + str(len(items)) + tp.group(3)
                    + "".join(items) + tp.group(5) + app[tp.end():]
                )
                parts["docProps/app.xml"] = app.encode("utf-8")
        except Exception:
            pass
    return path, parts


def save_info_familias(source: Any, campos: dict[str, str], competencias: dict[str, str]) -> bytes:
    payload = _read_bytes(source)
    src = zipfile.ZipFile(BytesIO(payload))
    path = _sheet_path(src)
    replacements: dict[str, bytes] = {}
    nueva = path is None
    if nueva:
        path, replacements = _registrar_hoja(src)
    replacements[path] = _sheet_xml(campos, competencias).encode("utf-8")

    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, replacements.get(item.filename, src.read(item.filename)))
        if nueva:
            dst.writestr(path, replacements[path])
    return out.getvalue()


# ---------------------------------------------------------------------------
# CÁLCULOS
# ---------------------------------------------------------------------------
def pesos_il(il_rows: list[dict], ce_p: dict[str, float]) -> dict[str, float]:
    """Peso de cada IL sobre la calificación de la materia:
    ``(PIL / ΣPIL del CE) · (P / ΣP)``."""
    total_p = sum(ce_p.values())
    sum_pil: dict[str, float] = {}
    for r in il_rows:
        sum_pil[_s(r.get("CE"))] = sum_pil.get(_s(r.get("CE")), 0.0) + float(r.get("PIL") or 0)
    out = {}
    for r in il_rows:
        ce = _s(r.get("CE"))
        spil = sum_pil.get(ce, 0.0)
        out[_s(r.get("IL"))] = (
            (float(r.get("PIL") or 0) / spil if spil else 0.0)
            * (ce_p.get(ce, 0.0) / total_p if total_p else 0.0)
        )
    return out


def _tokens(value: Any) -> list[str]:
    return [re.sub(r"\s+", "", t) for t in str(value or "").split(",") if t.strip()]


def datos_familias(
    *,
    sits: list[dict],
    ce_list: list[str],
    ce_ced: dict[str, str],
    ce_p: dict[str, float],
    il_rows: list[dict],
    elementos: dict,
) -> dict:
    """Todo lo calculado para el documento (y la vista previa de la app).

    ``sits``: [{"SA", "EV", "DSA", "HSA"}] de la tabla de situaciones (la que
    manda). ``il_rows``: indicadores ya recalculados (IL/PIL% rellenos)."""
    total_p = sum(ce_p.values())
    pct_ce = {ce: (ce_p.get(ce, 0.0) / total_p if total_p else 0.0) for ce in ce_list}
    w_il = pesos_il(il_rows, ce_p)

    sits = sorted(
        (s for s in sits if s.get("SA") not in (None, "")), key=lambda s: int(s["SA"])
    )
    total_h = sum(float(s.get("HSA") or 0) for s in sits)
    ev_por_sa = {int(s["SA"]): s.get("EV") for s in sits}

    # ── Criterios: peso, instrumentos y evaluaciones en que se trabajan
    criterios = []
    for ce in ce_list:
        ils = [r for r in il_rows if _s(r.get("CE")) == ce]
        ies = list(dict.fromkeys(_s(r.get("IE")) for r in ils if _s(r.get("IE"))))
        evs = sorted({
            int(ev_por_sa[int(r["SA"])]) for r in ils
            if r.get("SA") not in (None, "") and ev_por_sa.get(int(r["SA"])) not in (None, "")
        })
        criterios.append({
            "CE": ce, "CED": ce_ced.get(ce, ""), "pct": pct_ce.get(ce, 0.0),
            "IE": ", ".join(ies), "EV": ", ".join(f"{e}ª" for e in evs),
        })

    # ── Competencias específicas presentes (primer nº de cada CE) y su peso
    comp_peso: dict[str, float] = {}
    for ce in ce_list:
        n = numero_competencia(ce)
        comp_peso[n] = comp_peso.get(n, 0.0) + pct_ce.get(ce, 0.0)

    # ── Contenidos: en qué SA/evaluación se trabaja cada uno
    con_sa: dict[str, set[int]] = {}
    for r in il_rows:
        if r.get("SA") in (None, ""):
            continue
        for tok in _tokens(r.get("CON")):
            con_sa.setdefault(tok, set()).add(int(r["SA"]))
    contenidos = []
    for it in elementos.get("contenidos", []):
        sas = sorted(con_sa.get(re.sub(r"\s+", "", it["cod"]), set()))
        evs = sorted({int(ev_por_sa[s]) for s in sas if ev_por_sa.get(s) not in (None, "")})
        contenidos.append({
            "cod": it["cod"], "desc": it.get("desc", ""),
            "SA": ", ".join(str(s) for s in sas), "EV": ", ".join(f"{e}ª" for e in evs),
        })

    # ── Temporalización
    temporalizacion = [
        {
            "EV": s.get("EV"), "SA": int(s["SA"]), "DSA": _s(s.get("DSA")),
            "HSA": s.get("HSA"),
            "pct_h": (float(s.get("HSA") or 0) / total_h) if total_h else 0.0,
        }
        for s in sorted(sits, key=lambda s: (int(s.get("EV") or 9), int(s["SA"])))
    ]

    # ── Instrumentos (global)
    por_ie: dict[str, float] = {}
    for r in il_rows:
        if _s(r.get("IE")):
            por_ie[_s(r["IE"])] = por_ie.get(_s(r["IE"]), 0.0) + w_il.get(_s(r.get("IL")), 0.0)
    instrumentos = sorted(por_ie.items(), key=lambda kv: -kv[1])

    # ── Una entrada por SA: criterios → indicadores con sus pesos
    situaciones = []
    for s in sits:
        n = int(s["SA"])
        ils = [r for r in il_rows if r.get("SA") not in (None, "") and int(r["SA"]) == n]
        peso_sa = sum(w_il.get(_s(r.get("IL")), 0.0) for r in ils)
        grupos = []
        for ce in list(dict.fromkeys(_s(r.get("CE")) for r in ils)):
            filas = []
            for r in (r for r in ils if _s(r.get("CE")) == ce):
                w = w_il.get(_s(r.get("IL")), 0.0)
                filas.append({
                    "IL": _s(r.get("IL")),
                    "pil_pct": float(r.get("PIL%") or 0.0),
                    "peso_sa": (w / peso_sa) if peso_sa else 0.0,
                    "peso_materia": w,
                    "IE": _s(r.get("IE")), "CC": _s(r.get("CC")), "AE": _s(r.get("AE")),
                    "CON": ", ".join(_tokens(r.get("CON"))), "CT": ", ".join(_tokens(r.get("CT"))),
                })
            grupos.append({
                "CE": ce, "CED": ce_ced.get(ce, ""), "pct_ce": pct_ce.get(ce, 0.0),
                "peso_sa": sum(f["peso_sa"] for f in filas), "filas": filas,
            })
        ie_sa: dict[str, float] = {}
        for g in grupos:
            for f in g["filas"]:
                if f["IE"]:
                    ie_sa[f["IE"]] = ie_sa.get(f["IE"], 0.0) + f["peso_sa"]
        situaciones.append({
            "numero": n, "titulo": _s(s.get("DSA")),
            "EV": s.get("EV"), "HSA": s.get("HSA"),
            "pct_h": (float(s.get("HSA") or 0) / total_h) if total_h else 0.0,
            "peso": peso_sa, "grupos": grupos,
            "instrumentos": sorted(ie_sa.items(), key=lambda kv: -kv[1]),
            "agentes": sorted({f["AE"] for g in grupos for f in g["filas"] if f["AE"]}),
        })

    # ── Resumen por criterio: el mismo que va a la programación (hoja RESUMEN)
    from tools.indicadores_logro import resumen_por_ce

    resumen_rows = [
        [r["CE"], _fmt_pct(r["pct"]), r["CONTENIDOS"], r["CT"], r["SA"]]
        for r in resumen_por_ce(il_rows, ce_list, ce_ced, ce_p)
    ]

    return {
        "resumen_rows": resumen_rows,
        "criterios": criterios,
        "competencias": sorted(comp_peso.items(), key=lambda kv: _comp_sort_key(kv[0])),
        "contenidos": contenidos,
        "temporalizacion": temporalizacion,
        "instrumentos": instrumentos,
        "situaciones": situaciones,
    }


def tabla_sa_filas(sa: dict) -> tuple[list[str], list[list[str]]]:
    """Tabla plana de una SA (para la vista previa de la app)."""
    headers = ["CE / IL", "Peso IL en el CE", "Peso en la SA", "Peso en la materia",
               "Instrumento", "Calificación", "Agente", "Contenidos", "C. transversales"]
    rows = []
    for g in sa["grupos"]:
        rows.append([f"CE {g['CE']}", "", _fmt_pct(g["peso_sa"]), "", "", "", "", "", ""])
        for f in g["filas"]:
            rows.append([
                f"   {f['IL']}", _fmt_pct(f["pil_pct"]), _fmt_pct(f["peso_sa"]),
                _fmt_pct(f["peso_materia"]), f["IE"], f["CC"], f["AE"], f["CON"], f["CT"],
            ])
    return headers, rows


# ---------------------------------------------------------------------------
# WORD
# ---------------------------------------------------------------------------
DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "informacion_familias_template.docx"


def _font(cell, size: float, *, bold: bool = False, color: RGBColor | None = None) -> None:
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold or run.font.bold
            if color is not None:
                run.font.color.rgb = color


def _fijar_anchos(table, widths_cm: list[float]) -> None:
    """Word ignora los anchos de celda si la tabla está en autoajuste: se fija
    el diseño y la rejilla de columnas, además de cada celda."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for col, w in zip(table.columns, widths_cm):
        col.width = Cm(w)
    for row in table.rows:
        cells = row.cells
        if len({id(c._tc) for c in cells}) < len(cells):
            continue  # fila con celdas combinadas (cabecera de criterio)
        for cell, w in zip(cells, widths_cm):
            cell.width = Cm(w)


MAX_FILAS_JUNTAS = 20  # tablas hasta este tamaño no se parten entre páginas


def _fila(row, *, cabecera: bool = False, con_siguiente: bool = False) -> None:
    """Maquetación de una fila: nunca se parte entre dos páginas; la cabecera
    se repite en cada página; ``con_siguiente`` la ata a la fila (o párrafo)
    que viene después, para que Word no corte ahí."""
    from docx.oxml import OxmlElement

    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    if cabecera:
        trPr.append(OxmlElement("w:tblHeader"))
    if con_siguiente or cabecera:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


def _tabla(tpl: DocxTemplate, headers: list[str], rows: list[list[Any]], widths_cm: list[float],
           *, size: float = 9, empty_msg: str = "Sin datos.", con_lo_siguiente: bool = False,
           max_juntas: int = MAX_FILAS_JUNTAS):
    """``con_lo_siguiente``: la tabla (si es corta) se mantiene en la misma
    página que el párrafo de después (p. ej. su gráfico). ``max_juntas``: hasta
    cuántas filas se mantiene entera (menos en tablas de filas altas)."""
    sub = tpl.new_subdoc()
    if not rows:
        sub.add_paragraph().add_run(empty_msg).italic = True
        return sub
    t = sub.add_table(rows=1, cols=len(headers))
    _set_table_borders(t)
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        c.width = Cm(widths_cm[i])
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(c, size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _set_cell_background(c, HEADER_FILL)
    _fila(t.rows[0], cabecera=True)
    juntas = len(rows) <= max_juntas
    total = len(rows)
    for n, vals in enumerate(rows, 1):
        row = t.add_row()
        cells = row.cells
        for i, v in enumerate(vals):
            cells[i].text = "" if v is None else str(v)
            cells[i].width = Cm(widths_cm[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _font(cells[i], size)
        if juntas:
            atar = n < total or con_lo_siguiente
        else:  # tabla larga: que no empiece ni acabe con filas sueltas
            atar = n <= 2 or total - 3 < n < total
        _fila(row, con_siguiente=atar)
    _fijar_anchos(t, widths_cm)
    return sub


def _tabla_sa(tpl: DocxTemplate, sa: dict):
    """Criterio (fila destacada con su descripción) y debajo sus indicadores
    de logro con todos los valores, sin el texto del indicador."""
    sub = tpl.new_subdoc()
    if not sa["grupos"]:
        sub.add_paragraph().add_run(
            "No hay indicadores de logro asociados a esta situación de aprendizaje."
        ).italic = True
        return sub
    headers = ["Indicador", "Peso en el criterio", "Peso en la SA", "Peso en la materia",
               "Instrumento", "Calificación", "Agente", "Contenidos", "C. transv."]
    widths = [1.5, 1.6, 1.5, 1.6, 2.3, 1.9, 1.4, 3.4, 1.8]
    t = sub.add_table(rows=1, cols=len(headers))
    _set_table_borders(t)
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        c.width = Cm(widths[i])
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(c, 8, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _set_cell_background(c, HEADER_FILL)
    _fila(t.rows[0], cabecera=True)
    for gi, g in enumerate(sa["grupos"]):
        ultimo_grupo = gi == len(sa["grupos"]) - 1
        row = t.add_row()
        cells = row.cells
        m = cells[0].merge(cells[-1])
        m.text = ""
        p = m.paragraphs[0]
        r = p.add_run(f"CE {g['CE']}. ")
        r.bold = True
        r.font.size = Pt(8.5)
        r = p.add_run(g["CED"])
        r.font.size = Pt(8.5)
        r = p.add_run(
            f"  (peso del criterio en la materia: {_fmt_pct(g['pct_ce'])} · "
            f"en esta SA: {_fmt_pct(g['peso_sa'])})"
        )
        r.italic = True
        r.font.size = Pt(8)
        _set_cell_background(m, "FAF4EC")
        # después de escribir el texto (``cell.text`` borra el formato del párrafo)
        _fila(row, con_siguiente=True)  # el criterio no se separa de sus indicadores
        for fi, f in enumerate(g["filas"]):
            vals = [f["IL"], _fmt_pct(f["pil_pct"]), _fmt_pct(f["peso_sa"]),
                    _fmt_pct(f["peso_materia"]), f["IE"], f["CC"], f["AE"], f["CON"], f["CT"]]
            row = t.add_row()
            cells = row.cells
            for i, v in enumerate(vals):
                cells[i].text = v
                cells[i].width = Cm(widths[i])
                cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if 0 < i < 4 or i == 6:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _font(cells[i], 8)
            # Dentro del criterio todo va junto; el último criterio de la SA se
            # ata además a la línea de agentes y al gráfico (que no quede solo).
            _fila(row, con_siguiente=fi < len(g["filas"]) - 1 or ultimo_grupo)
    _fijar_anchos(t, widths)
    return sub


def _texto(valor: str) -> RichText:
    rt = RichText()
    rt.add(valor.strip() or SIN_COMPLETAR, italic=not valor.strip())
    return rt


def campos_vacios(campos: dict[str, str], competencias: dict[str, str], nums: list[str]) -> list[str]:
    """Etiquetas de lo que falta por rellenar (para avisar antes de generar)."""
    faltan = [lbl for k, lbl, _h in CAMPOS if not _s(campos.get(k))]
    faltan += [f"Competencia específica {n}" for n in nums if not _s(competencias.get(n))]
    return faltan


def generar_docx(
    *,
    datos_generales: dict[str, str],
    calculo: dict,
    campos: dict[str, str],
    competencias: dict[str, str],
    template_source: Any = None,
) -> bytes:
    from tools.situaciones_informe import build_pie_png_simple

    if template_source is not None:
        tpl_bytes = _read_bytes(template_source)
    elif DEFAULT_TEMPLATE.exists():
        tpl_bytes = DEFAULT_TEMPLATE.read_bytes()
    else:
        from tools.plantilla_familias import build_template_bytes

        tpl_bytes = build_template_bytes()
    tpl = DocxTemplate(BytesIO(tpl_bytes))

    def _img(labels, values, titulo, ancho=12.0):
        png = build_pie_png_simple(labels, values, titulo=titulo, leyenda="Instrumento")
        return InlineImage(tpl, BytesIO(png), width=Cm(ancho))

    ctx: dict[str, Any] = {k: _s(v) for k, v in datos_generales.items()}
    ctx["anio"] = _s(datos_generales.get("años"))

    ctx["tabla_competencias"] = _tabla(
        tpl, ["Nº", "Competencia específica", "Peso en la materia"],
        [[n, _s(competencias.get(n)) or SIN_COMPLETAR, _fmt_pct(p)] for n, p in calculo["competencias"]],
        [1.2, 13.2, 2.6],
    )
    ctx["tabla_criterios"] = _tabla(
        tpl, ["CE", "Criterio de evaluación", "Peso", "Instrumentos de evaluación", "Evaluación"],
        [[c["CE"], c["CED"], _fmt_pct(c["pct"]), c["IE"], c["EV"]] for c in calculo["criterios"]],
        [1.1, 8.8, 1.7, 3.6, 1.8], size=8.5, max_juntas=6,
    )
    ctx["tabla_resumen"] = _tabla(
        tpl, ["CE", "%CE", "CONTENIDOS", "CT", "SA"],
        calculo["resumen_rows"], [1.2, 1.8, 9.0, 2.5, 2.5], size=8.5, max_juntas=10,
    )
    ctx["tabla_contenidos"] = _tabla(
        tpl, ["Código", "Contenido", "Situaciones", "Evaluación"],
        [[c["cod"], c["desc"], c["SA"] or "—", c["EV"] or "—"] for c in calculo["contenidos"]],
        [1.6, 11.0, 2.4, 2.0], size=8, empty_msg="No hay contenidos en la hoja LOMLOE.", max_juntas=10,
    )
    ctx["tabla_temporalizacion"] = _tabla(
        tpl, ["Evaluación", "SA", "Situación de aprendizaje", "Horas", "% horas"],
        [[f"{t['EV']}ª" if t["EV"] else "—", t["SA"], t["DSA"],
          "" if t["HSA"] in (None, "") else f"{t['HSA']:g}", _fmt_pct(t["pct_h"])]
         for t in calculo["temporalizacion"]],
        [2.0, 1.2, 9.4, 1.8, 2.6],
    )
    ins = calculo["instrumentos"]
    ctx["tabla_instrumentos"] = _tabla(
        tpl, ["Instrumento de evaluación", "Peso en la calificación"],
        [[k, _fmt_pct(v)] for k, v in ins] + ([["Total", _fmt_pct(sum(v for _, v in ins))]] if ins else []),
        [10.0, 5.0], con_lo_siguiente=True,
    )
    ctx["grafico_instrumentos"] = (
        _img([k for k, _ in ins], [v for _, v in ins], "Peso por instrumento de evaluación")
        if ins else ""
    )
    for k in CAMPOS_KEYS:
        ctx[k] = _texto(_s(campos.get(k)))

    situaciones = []
    for sa in calculo["situaciones"]:
        situaciones.append({
            "numero": sa["numero"],
            "titulo": sa["titulo"],
            "trimestre": f"{sa['EV']}ª evaluación" if sa["EV"] else "",
            "horas": "" if sa["HSA"] in (None, "") else f"{sa['HSA']:g}",
            "pct_horas": _fmt_pct(sa["pct_h"]),
            "peso": _fmt_pct(sa["peso"]),
            "tabla": _tabla_sa(tpl, sa),
            "grafico": (
                _img([k for k, _ in sa["instrumentos"]], [v for _, v in sa["instrumentos"]],
                     f"Instrumentos de evaluación · SA {sa['numero']}", ancho=10.5)
                if sa["instrumentos"] else ""
            ),
            "agentes": "; ".join(f"{a} = {AGENTES.get(a, a)}" for a in sa["agentes"]),
        })
    ctx["situaciones"] = situaciones

    tpl.render(ctx, jinja_env=_JINJA_ENV)
    out = BytesIO()
    tpl.save(out)
    return out.getvalue()
