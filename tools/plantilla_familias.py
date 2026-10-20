"""Construye la plantilla Word (docxtpl) por defecto del documento de
información a las familias: ``templates/informacion_familias_template.docx``.

Se puede retocar luego a mano en Word (estilos, cabecera del centro…) siempre
que se respeten las etiquetas ``{{ … }}`` / ``{%p … %}``. Para regenerarla:

    .venv\\Scripts\\python.exe -m tools.plantilla_familias
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROJO = RGBColor(0xA6, 0x18, 0x2E)
TINTA = RGBColor(0x22, 0x1C, 0x18)

INTRO = (
    "De acuerdo con la ORDEN EDU/424/2024, de 9 de mayo (art. 12, ESO) y la ORDEN "
    "EDU/425/2024, de 9 de mayo (art. 10, Bachillerato), al principio de cada curso "
    "el profesorado informa al alumnado y a las familias de las competencias "
    "específicas, los criterios de evaluación, los contenidos, los procedimientos e "
    "instrumentos de evaluación y los criterios de calificación de la materia, así "
    "como de las medidas de atención individualizada. Esta información procede de la "
    "programación didáctica del departamento."
)


def _estilos(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TINTA
    for nombre, tam in (("Title", 22), ("Heading 1", 14), ("Heading 2", 13)):
        st = doc.styles[nombre]
        st.font.name = "Calibri"
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = ROJO
        st.paragraph_format.keep_with_next = True
    for sec in doc.sections:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(1.8)


def _p(doc: Document, texto: str = "", *, bold: bool = False, italic: bool = False,
       center: bool = False, size: float | None = None, junto: bool = False):
    """``junto``: el párrafo no se separa del siguiente (texto que presenta una
    tabla o un gráfico)."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = junto
    if texto:
        r = p.add_run(texto)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _etiquetado(doc: Document, etiqueta: str, var: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{etiqueta}: ").bold = True
    p.add_run(f"{{{{r {var} }}}}")


def build_document() -> Document:
    doc = Document()
    _estilos(doc)

    doc.add_paragraph("Información a las familias sobre la evaluación", style="Title")
    _p(doc, "{{ materia }} · {{ curso }} {{ grupo }} · Curso {{ anio }}", bold=True, size=12)
    _p(doc, "Docente: {{ docente }}")
    _p(doc, INTRO, italic=True, size=9.5)

    doc.add_heading("1. Competencias específicas", level=1)
    _p(doc, "{{p tabla_competencias }}")

    doc.add_heading("2. Criterios de evaluación y su peso en la calificación", level=1)
    _p(doc, "El peso indica cuánto cuenta cada criterio en la calificación final de la materia.",
       size=9.5, junto=True)
    _p(doc, "{{p tabla_criterios }}")

    doc.add_heading("3. Resumen por criterio de evaluación", level=1)
    _p(doc, "Contenidos (CON), contenidos transversales (CT) y situaciones de aprendizaje (SA) "
       "asociados a cada criterio.", size=9.5, junto=True)
    _p(doc, "{{p tabla_resumen }}")

    doc.add_heading("4. Contenidos", level=1)
    _p(doc, "{{p tabla_contenidos }}")

    doc.add_heading("5. Organización de los contenidos a lo largo del curso", level=1)
    _p(doc, "Secuenciación y temporalización de las situaciones de aprendizaje.", size=9.5,
       junto=True)
    _p(doc, "{{p tabla_temporalizacion }}")

    doc.add_heading("6. Procedimientos e instrumentos de evaluación", level=1)
    _p(doc, "{{p tabla_instrumentos }}")
    _p(doc, "{{ grafico_instrumentos }}", center=True)

    doc.add_heading("7. Criterios de calificación", level=1)
    _etiquetado(doc, "Redondeo de la calificación de cada criterio", "redondeo_criterios")
    _etiquetado(doc, "Redondeo de la calificación de la materia", "redondeo_materia")
    _etiquetado(doc, "Competencias clave", "competencias_clave")

    doc.add_heading("8. Criterios de promoción y titulación", level=1)
    _p(doc, "{{r promocion_titulacion }}")

    doc.add_heading("9. Medidas de atención individualizada", level=1)
    _etiquetado(doc, "Plan de recuperación", "plan_recuperacion")
    _etiquetado(doc, "Plan de refuerzo", "plan_refuerzo")
    _etiquetado(doc, "Apoyos individualizados o en grupo específico", "apoyos")
    _etiquetado(doc, "Adaptaciones curriculares", "adaptaciones")

    # Una página por situación de aprendizaje (salto antes del título: así no
    # queda un párrafo vacío arriba de la página).
    _p(doc, "{%p for sa in situaciones %}")
    h = doc.add_heading("Situación de aprendizaje {{ sa.numero }}. {{ sa.titulo }}", level=2)
    h.paragraph_format.page_break_before = True
    _p(doc, "{{ sa.trimestre }} · {{ sa.horas }} h ({{ sa.pct_horas }} de las horas del curso) "
       "· Peso en la calificación de la materia: {{ sa.peso }}", bold=True, size=9.5, junto=True)
    _p(doc, "Criterios de evaluación que se trabajan y sus indicadores de logro.", size=9.5,
       junto=True)
    _p(doc, "{{p sa.tabla }}")
    _p(doc, "Agente evaluador: {{ sa.agentes }}", italic=True, size=8.5, junto=True)
    _p(doc, "{{ sa.grafico }}", center=True)
    _p(doc, "{%p endfor %}")
    return doc


def build_template_bytes() -> bytes:
    buf = BytesIO()
    build_document().save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    destino = Path(__file__).resolve().parent.parent / "templates" / "informacion_familias_template.docx"
    destino.write_bytes(build_template_bytes())
    print(f"Plantilla escrita en {destino}")
