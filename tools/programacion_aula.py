from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import jinja2
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docxtpl import DocxTemplate

# Un índice/TOC guardado en la plantilla puede dejar en caché texto de una versión
# anterior de un encabezado (p.ej. "{{ situacion.titulo }}") fuera del bucle real.
# Con el Undefined por defecto de Jinja eso rompe el render ('situacion' is undefined);
# con ChainableUndefined esas referencias sueltas se imprimen en blanco sin fallar.
_JINJA_ENV = jinja2.Environment(undefined=jinja2.ChainableUndefined)

INDICADORES_COLUMNAS = ["CE", "IL", "CT", "IE", "CC", "AE"]
INDICADORES_CABECERAS = INDICADORES_COLUMNAS

TABLE_HEADER_FILL = "000000"
TABLE_HEADER_FONT_SIZE_PT = 9


def _read_source_bytes(source: Any) -> bytes:
    if isinstance(source, (str, Path)):
        path = Path(source).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"No existe el archivo: {path}")
        return path.read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    raise TypeError("El origen debe ser una ruta o un archivo subido")


def ensure_template(template_source: Any) -> bytes:
    if template_source is not None:
        return _read_source_bytes(template_source)

    default_template = Path(__file__).resolve().parent.parent / "templates" / "programacion_aula_template.docx"
    if not default_template.exists():
        raise FileNotFoundError(
            f"No se ha subido plantilla y no existe la plantilla por defecto: {default_template}"
        )
    return default_template.read_bytes()


def _fmt_pct(value: float) -> str:
    return f"{value * 100:.2f}".replace(".", ",") + " %"


def _set_cell_background(cell, color_hex: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _set_table_borders(table) -> None:
    """Bordes finos vía XML directo: no depende de que exista el estilo con
    nombre 'Table Grid' en la plantilla (puede faltar si Word nunca lo llegó
    a aplicar y lo optimizó fuera de styles.xml al guardar)."""
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tblPr.append(borders)


def _build_table_subdoc(tpl: DocxTemplate, headers: list[str], rows: list[list[str]], empty_msg: str):
    subdoc = tpl.new_subdoc()

    if not rows:
        p = subdoc.add_paragraph()
        run = p.add_run(empty_msg)
        run.italic = True
        return subdoc

    table = subdoc.add_table(rows=1, cols=len(headers))
    _set_table_borders(table)
    table.autofit = True

    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = header_cells[i].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(TABLE_HEADER_FONT_SIZE_PT)
        _set_cell_background(header_cells[i], TABLE_HEADER_FILL)

    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = "" if value is None else str(value)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    return subdoc


def build_tabla_indicadores(tpl: DocxTemplate, indicadores: list[dict]):
    rows = [[fila.get(col) for col in INDICADORES_COLUMNAS] for fila in indicadores]
    return _build_table_subdoc(
        tpl,
        INDICADORES_CABECERAS,
        rows,
        empty_msg="No hay indicadores de logro asociados a esta situación de aprendizaje.",
    )


def _norm_titulo(value: Any) -> str:
    import re as _re

    return _re.sub(r"\s+", " ", str(value or "")).strip().lower()


def build_tabla_actividades(tpl: DocxTemplate, actividades: list[dict]):
    headers = ["A", "Descripción", "Valor s/ programación", "Valor s/ SA"]
    rows = [
        [a.get("A"), a.get("DA"), _fmt_pct(a.get("valor_programacion", 0.0)), _fmt_pct(a.get("valor_sa", 0.0))]
        for a in actividades
    ]
    return _build_table_subdoc(
        tpl,
        headers,
        rows,
        empty_msg="No hay actividades asociadas a esta situación de aprendizaje.",
    )


def run_programacion_aula(
    datos_generales: dict,
    situaciones: list[dict],
    template_source: Any = None,
    *,
    resumen_por_sa: dict | None = None,
) -> bytes:
    """Genera el .docx de la programación de aula.

    ``datos_generales``: ``{campo: valor}`` (variables sueltas de la plantilla).
    ``situaciones``: una por SA, en orden, con ``numero``, ``titulo``,
    ``trimestre`` y los campos de aula (``contextualizacion``, ``metodos``…).
    ``resumen_por_sa``: ``{titulo_de_SA: {"indicadores": [...], "actividades":
    [...]}}`` ya calculado por la app (misma cuenta que el «Resumen de
    actividades» de cada SA en pantalla) para las tablas
    ``{{p situacion.tabla_indicadores}}`` y ``{{p situacion.tabla_actividades}}``.
    """
    if not situaciones:
        raise ValueError("No hay situaciones de aprendizaje.")

    tpl = DocxTemplate(BytesIO(ensure_template(template_source)))
    resumen = {_norm_titulo(k): v for k, v in (resumen_por_sa or {}).items()}

    ctx = dict(datos_generales)
    ctx["situaciones"] = []
    for sit in situaciones:
        sit = dict(sit)
        entrada = resumen.get(_norm_titulo(sit.get("titulo")), {})
        sit["tabla_indicadores"] = build_tabla_indicadores(tpl, entrada.get("indicadores", []))
        sit["tabla_actividades"] = build_tabla_actividades(tpl, entrada.get("actividades", []))
        ctx["situaciones"].append(sit)

    tpl.render(ctx, jinja_env=_JINJA_ENV)
    buffer = BytesIO()
    tpl.save(buffer)
    return buffer.getvalue()
