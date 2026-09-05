"""Lectura y escritura de la tabla `tablaSAprendizaje` (hoja SituacionesAprendizaje).

El Excel de programación lleva un modelo de Power Pivot, tablas dinámicas,
segmentaciones y datos enriquecidos. `openpyxl` (y `pandas`) al guardar tiran
todo eso a la basura (el archivo pasa de ~690 KB a ~60 KB). Por eso aquí:

* la LECTURA se hace con openpyxl (solo lectura, no daña nada);
* la ESCRITURA es quirúrgica: se abre el .xlsx como ZIP y se reescriben
  únicamente las partes mínimas afectadas
  (`xl/worksheets/<hoja SA>.xml`, la definición de su tabla y `xl/workbook.xml`),
  copiando **byte a byte** todo lo demás. Así el modelo de datos, las dinámicas
  y el formato condicional quedan intactos y Excel los recalcula al abrir.

Estructura de la hoja:

* Tabla `tablaSAprendizaje`, rango `A1:D{n}`, columnas ``SA``, ``EV``, ``DSA``,
  ``HSA`` (nº de situación, evaluación/trimestre 1-3, descripción, horas).
* Bloque de celdas suelto ``G1:J4`` (no es una tabla) con el resumen por
  trimestre: ``EVALUACIÓN`` (1-3), ``HORAS PREVISTAS`` (a mano),
  ``ACUMULADO`` (``=SUMIF(tablaSAprendizaje[EV];Gn;tablaSAprendizaje[HSA])``) y
  ``DESVIACIÓN`` (``=Hn-In``), con formato condicional que pinta la desviación
  de rojo si es negativa (te has pasado de horas) y de verde si es >= 0.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import load_workbook

SHEET_NAME = "SituacionesAprendizaje"
TABLE_NAME = "tablaSAprendizaje"
TABLE_COLUMNS = ["SA", "EV", "DSA", "HSA"]

# Bloque resumen G1:J4
RESUMEN_FIRST_ROW = 2
RESUMEN_LAST_ROW = 4
TRIMESTRES = (1, 2, 3)


@dataclass
class Situacion:
    sa: int | None
    ev: int | None
    dsa: str
    hsa: float | None


@dataclass
class SituacionesData:
    situaciones: list[Situacion]
    horas_previstas: dict[int, float]  # {trimestre: horas}


# ---------------------------------------------------------------------------
# Utilidades comunes
# ---------------------------------------------------------------------------
def _read_source_bytes(source: Any) -> bytes:
    if isinstance(source, (str, Path)):
        path = Path(source).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"No existe el archivo: {path}")
        return path.read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    if hasattr(source, "read"):
        return source.read()
    raise TypeError("El origen debe ser una ruta o un archivo subido")


def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        f = float(value)
    except (TypeError, ValueError):
        return None
    return int(f) if f.is_integer() else f


def _num_xml(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else repr(float(value))


# ---------------------------------------------------------------------------
# LECTURA (openpyxl, solo lectura)
# ---------------------------------------------------------------------------
def read_situaciones(source: Any) -> SituacionesData:
    payload = _read_source_bytes(source)
    wb = load_workbook(BytesIO(payload), data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"El Excel no tiene una hoja llamada '{SHEET_NAME}'.")
    ws = wb[SHEET_NAME]

    # Cabecera real de la tabla: fila que contiene 'SA' en la columna A.
    header_row = None
    for row in ws.iter_rows(min_col=1, max_col=1, max_row=20):
        if row[0].value is not None and str(row[0].value).strip() == "SA":
            header_row = row[0].row
            break
    if header_row is None:
        raise ValueError(
            f"No se encuentra la tabla '{TABLE_NAME}' en la hoja '{SHEET_NAME}'."
        )

    headers = [
        (str(c.value).strip() if c.value is not None else "")
        for c in ws[header_row][:4]
    ]
    idx = {name: headers.index(name) for name in TABLE_COLUMNS if name in headers}
    for needed in TABLE_COLUMNS:
        if needed not in idx:
            raise ValueError(
                f"La tabla '{TABLE_NAME}' no tiene la columna '{needed}'. "
                f"Columnas encontradas: {headers}"
            )

    situaciones: list[Situacion] = []
    for row in ws.iter_rows(min_row=header_row + 1, max_col=4):
        values = [c.value for c in row]
        if all(v is None or v == "" for v in values):
            continue
        situaciones.append(
            Situacion(
                sa=_num(values[idx["SA"]]),
                ev=_num(values[idx["EV"]]),
                dsa="" if values[idx["DSA"]] is None else str(values[idx["DSA"]]).strip(),
                hsa=_num(values[idx["HSA"]]),
            )
        )

    # Bloque resumen G:J -> horas previstas por trimestre (col G = nº, col H = horas).
    horas_previstas: dict[int, float] = {}
    for r in range(RESUMEN_FIRST_ROW, RESUMEN_LAST_ROW + 1):
        tri = _num(ws.cell(row=r, column=7).value)  # G
        horas = _num(ws.cell(row=r, column=8).value)  # H
        if tri in TRIMESTRES:
            horas_previstas[int(tri)] = horas if horas is not None else 0.0
    for tri in TRIMESTRES:
        horas_previstas.setdefault(tri, 0.0)

    return SituacionesData(situaciones=situaciones, horas_previstas=horas_previstas)


def acumulado_por_trimestre(situaciones: list[Situacion]) -> dict[int, float]:
    acc = {tri: 0.0 for tri in TRIMESTRES}
    for s in situaciones:
        if s.ev in TRIMESTRES and s.hsa:
            acc[int(s.ev)] += float(s.hsa)
    return acc


# ---------------------------------------------------------------------------
# ESCRITURA (quirúrgica sobre el ZIP)
# ---------------------------------------------------------------------------
_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _find_sheet_path(zf: zipfile.ZipFile) -> str:
    wb = zf.read("xl/workbook.xml").decode("utf-8")
    m = re.search(
        rf'<sheet[^>]*name="{re.escape(SHEET_NAME)}"[^>]*r:id="([^"]+)"', wb
    )
    if not m:
        raise ValueError(f"No se encuentra la hoja '{SHEET_NAME}' en el libro.")
    rid = m.group(1)
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    t = re.search(rf'<Relationship[^>]*Id="{re.escape(rid)}"[^>]*Target="([^"]+)"', rels)
    if not t:
        raise ValueError("No se resuelve la relación de la hoja de situaciones.")
    target = t.group(1).lstrip("/")
    if not target.startswith("xl/"):
        target = "xl/" + target.replace("../", "")
    return target


def _find_table_path(zf: zipfile.ZipFile) -> str | None:
    for name in zf.namelist():
        if name.startswith("xl/tables/") and name.endswith(".xml"):
            xml = zf.read(name).decode("utf-8")
            if f'name="{TABLE_NAME}"' in xml:
                return name
    return None


def _col_letter(ref: str) -> str:
    return re.match(r"([A-Z]+)\d+", ref).group(1)


def _build_data_cells(row_num: int, s: Situacion) -> str:
    cells = []
    if s.sa is not None and s.sa != "":
        cells.append(f'<c r="A{row_num}"><v>{_num_xml(s.sa)}</v></c>')
    if s.ev is not None and s.ev != "":
        cells.append(f'<c r="B{row_num}"><v>{_num_xml(s.ev)}</v></c>')
    if s.dsa:
        cells.append(
            f'<c r="C{row_num}" t="inlineStr"><is>'
            f'<t xml:space="preserve">{escape(str(s.dsa))}</t></is></c>'
        )
    if s.hsa is not None and s.hsa != "":
        cells.append(f'<c r="D{row_num}"><v>{_num_xml(s.hsa)}</v></c>')
    return "".join(cells)


def _rewrite_sheet_xml(
    sheet_xml: str,
    situaciones: list[Situacion],
    horas_previstas: dict[int, float],
) -> str:
    sd = re.search(r"<sheetData>(.*)</sheetData>", sheet_xml, re.S)
    if not sd:
        raise ValueError("La hoja de situaciones no tiene <sheetData>.")

    rows_xml = re.findall(r"<row [^>]*>.*?</row>|<row [^>]*/>", sd.group(1), re.S)
    row_by_num: dict[int, str] = {}
    for rx in rows_xml:
        rn = int(re.search(r'r="(\d+)"', rx).group(1))
        row_by_num[rn] = rx

    header_row = row_by_num.get(1)
    if header_row is None:
        raise ValueError("La hoja de situaciones no tiene fila de cabecera.")

    # Celdas no pertenecientes a la tabla (columnas E..J): el bloque resumen G:J.
    extras_by_row: dict[int, str] = {}
    for rn, rx in row_by_num.items():
        if rn == 1:
            continue
        extra = [
            c for c in re.findall(r"<c [^>]*?/>|<c [^>]*?>.*?</c>", rx, re.S)
            if _col_letter(re.search(r'r="([A-Z]+\d+)"', c).group(1)) >= "E"
        ]
        if extra:
            extras_by_row[rn] = "".join(extra)

    # Actualiza HORAS PREVISTAS (H2:H4) dentro de las celdas extra conservadas.
    for rn in range(RESUMEN_FIRST_ROW, RESUMEN_LAST_ROW + 1):
        tri = rn - RESUMEN_FIRST_ROW + 1
        if rn in extras_by_row and tri in horas_previstas:
            nuevo = _num_xml(horas_previstas[tri])
            extras_by_row[rn] = re.sub(
                r'(<c r="H%d"[^>]*>)<v>[^<]*</v>' % rn,
                r"\g<1><v>%s</v>" % nuevo,
                extras_by_row[rn],
            )
    # Quita los valores en caché de ACUMULADO/DESVIACIÓN para forzar recálculo.
    for rn in range(RESUMEN_FIRST_ROW, RESUMEN_LAST_ROW + 1):
        if rn in extras_by_row:
            extras_by_row[rn] = re.sub(
                r'(<c r="[IJ]%d"[^>]*>(?:<f[^>]*>.*?</f>|<f[^>]*/>))<v>[^<]*</v>' % rn,
                r"\g<1>",
                extras_by_row[rn],
            )

    n = len(situaciones)
    last_row = max(n + 1, RESUMEN_LAST_ROW)

    new_rows = [header_row]
    for r in range(2, last_row + 1):
        data = _build_data_cells(r, situaciones[r - 2]) if r - 2 < n else ""
        extras = extras_by_row.get(r, "")
        new_rows.append(
            f'<row r="{r}" spans="1:10" x14ac:dyDescent="0.25">{data}{extras}</row>'
        )

    sheet_xml = (
        sheet_xml[: sd.start()]
        + "<sheetData>" + "".join(new_rows) + "</sheetData>"
        + sheet_xml[sd.end():]
    )
    sheet_xml = re.sub(
        r'<dimension ref="[^"]*"/>', f'<dimension ref="A1:J{last_row}"/>', sheet_xml
    )
    return sheet_xml


def _rewrite_table_xml(table_xml: str, n: int) -> str:
    ref = f"A1:D{n + 1}"
    table_xml = re.sub(r'(<table[^>]*\sref=")[^"]*(")', rf"\g<1>{ref}\g<2>", table_xml)
    table_xml = re.sub(r'(<autoFilter\s+ref=")[^"]*(")', rf"\g<1>{ref}\g<2>", table_xml)
    return table_xml


def _force_full_recalc(workbook_xml: str) -> str:
    if "fullCalcOnLoad" in workbook_xml:
        return workbook_xml
    if "<calcPr" in workbook_xml:
        return re.sub(r"<calcPr\s+", '<calcPr fullCalcOnLoad="1" ', workbook_xml, count=1)
    return workbook_xml.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')


def save_situaciones(
    source: Any,
    situaciones: list[Situacion],
    horas_previstas: dict[int, float] | None = None,
) -> bytes:
    """Devuelve los bytes de un .xlsx idéntico al original salvo la tabla
    `tablaSAprendizaje` (y las HORAS PREVISTAS del bloque resumen)."""
    payload = _read_source_bytes(source)
    situaciones = [s for s in situaciones if not _is_empty(s)]
    horas_previstas = horas_previstas or {}

    src = zipfile.ZipFile(BytesIO(payload))
    sheet_path = _find_sheet_path(src)
    table_path = _find_table_path(src)

    new_sheet = _rewrite_sheet_xml(
        src.read(sheet_path).decode("utf-8"), situaciones, horas_previstas
    )
    new_workbook = _force_full_recalc(src.read("xl/workbook.xml").decode("utf-8"))
    new_table = (
        _rewrite_table_xml(src.read(table_path).decode("utf-8"), len(situaciones))
        if table_path
        else None
    )

    replacements = {
        sheet_path: new_sheet.encode("utf-8"),
        "xl/workbook.xml": new_workbook.encode("utf-8"),
    }
    if new_table is not None:
        replacements[table_path] = new_table.encode("utf-8")

    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = replacements.get(item.filename, src.read(item.filename))
            dst.writestr(item, data)
    return out.getvalue()


def _is_empty(s: Situacion) -> bool:
    return (
        (s.sa is None or s.sa == "")
        and (s.ev is None or s.ev == "")
        and not (s.dsa or "").strip()
        and (s.hsa is None or s.hsa == "")
    )
